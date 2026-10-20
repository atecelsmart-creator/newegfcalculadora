"""
Relatório por SKU (arquivo separado da planilha do marketplace): OK, REVISAR ou ABAIXO DO PMA, com o motivo.
Gera .xlsx (com aba de resumo) e um .csv ao lado.
"""
import csv
from collections import Counter
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

COLUNAS = [
    ("sku", "SKU"), ("status", "Status"), ("motivo", "Motivo"), ("cenario", "Cenário"),
    ("codigo_fabricante", "Cód. fabricante"), ("ncm", "NCM"), ("aliq_icms", "Aliq ICMS (%)"),
    ("unitario", "Unitário"), ("perc_ipi", "IPI"), ("perc_frete", "Frete"),
    ("st_status", "ST (Econet)"), ("mva", "MVA"), ("confaz_status", "Confaz"),
    ("preco_cheio", "Preço cheio"), ("preco_promocional", "Preço promocional"),
    ("margem_promocional", "Margem promocional"), ("preco_enviado", "Preço enviado à planilha"),
    ("gravado", "Gravado"), ("hora", "Hora"),
]
# Só aparecem no relatório de um lote do Mercado Livre (nos outros não existem)
COLUNAS_ML = [
    ("ml_comissao", "Comissão ML"), ("ml_peso", "Peso (frete ML)"), ("ml_faixa_valor", "Faixa informada (frete ML)"),
    ("frete_ml", "Frete ML informado (R$)"), ("frete_ml_aplicado", "Frete ML aplicado (R$)"),
    ("faixa_frete_aplicada", "Faixa aplicada (frete ML)"),
]
PERCENTUAIS = {"perc_ipi", "perc_frete", "mva", "margem_promocional", "ml_comissao"}
CORES = {"OK": "C6EFCE", "REVISAR": "FFEB9C", "ABAIXO DO PMA": "F8CBAD"}
LARGURAS = {"motivo": 60, "sku": 18, "cenario": 22}


def _valor(chave: str, v):
    if chave == "gravado":
        return "Sim" if v else "Não"
    return v


def gerar(linhas: list[dict], pasta: Path, prefixo: str = "shopee_sp") -> Path:
    pasta.mkdir(parents=True, exist_ok=True)
    carimbo = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = pasta / f"relatorio_{prefixo}_{carimbo}.xlsx"

    # as colunas do ML entram logo antes de "Gravado" e só se o lote for do Mercado Livre
    colunas = list(COLUNAS)
    if any(l.get(k) is not None for l in linhas for k, _ in COLUNAS_ML):
        i = next(n for n, (k, _) in enumerate(colunas) if k == "gravado")
        colunas[i:i] = COLUNAS_ML

    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.append([t for _, t in colunas])
    for c in ws[1]:
        c.font = Font(bold=True)
    for linha in linhas:
        ws.append([_valor(k, linha.get(k)) for k, _ in colunas])
        cor = CORES.get(linha.get("status"))
        if cor:
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=cor)
        for j, (k, _) in enumerate(colunas, start=1):
            if k in PERCENTUAIS and isinstance(linha.get(k), (int, float)):
                ws.cell(ws.max_row, j).number_format = "0.00%"
    for j, (k, _) in enumerate(colunas, start=1):
        ws.column_dimensions[get_column_letter(j)].width = LARGURAS.get(k, 16)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Status", "Quantidade"])
    for c in resumo[1]:
        c.font = Font(bold=True)
    contagem = Counter(l.get("status") for l in linhas)
    for status in ("OK", "ABAIXO DO PMA", "REVISAR"):
        resumo.append([status, contagem.get(status, 0)])
    resumo.append(["Total", len(linhas)])
    resumo.column_dimensions["A"].width = 18
    wb.save(caminho)

    with open(caminho.with_suffix(".csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow([t for _, t in colunas])
        for linha in linhas:
            w.writerow([_valor(k, linha.get(k)) if linha.get(k) is not None else "" for k, _ in colunas])
    return caminho
