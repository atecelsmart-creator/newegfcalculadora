"""
Tabela PMA (Preço Sugerido Revenda).

O usuário sobe uma planilha Excel (.xlsx, .xlsm, .xlsb ou .xls). Este módulo lê
todas as abas que tenham as colunas CÓDIGO e REVENDA (e, se existir, DESCRIÇÃO
DO PRODUTO), guarda os itens num JSON em backend/data/ e responde consultas por
código — sem precisar reler o Excel a cada cálculo.

Subir uma nova planilha SUBSTITUI a anterior. A comparação "preço promocional
abaixo do PMA" é feita em avaliar(), usada pela rota /calcular.
"""
import io
import json
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from threading import Lock

ARQUIVO_JSON = Path(__file__).resolve().parents[2] / "data" / "pma_tabela.json"
EXTENSOES = (".xlsx", ".xlsm", ".xlsb", ".xls")
LINHAS_PARA_ACHAR_CABECALHO = 30

# Nome da aba da planilha PMA -> marca dos produtos daquela aba (cada aba tem a sua marca).
# Aba que não estiver aqui mostra o próprio nome da aba. A chave é normalizada (maiúsculo, sem espaços).
MARCAS_POR_ABA = {
    "B+D": "Black + Decker",
    "B&D": "Black + Decker",
    "STANLEY": "Stanley",
    "DEWALT": "Dewalt",
}


def marca_da_aba(aba: str | None) -> str | None:
    if not aba:
        return None
    chave = re.sub(r"\s+", "", str(aba)).upper()
    return MARCAS_POR_ABA.get(chave, str(aba).strip().title())

_lock = Lock()
_tabela: dict = {"arquivo": None, "carregada_em": None, "itens": {}}


# ---------------------------------------------------------------- normalização
def _texto_chave(valor) -> str:
    """Maiúsculo, sem acento e sem espaços repetidos (pra reconhecer cabeçalhos)."""
    s = unicodedata.normalize("NFKD", str(valor))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", s).strip().upper()


def normalizar_codigo(valor) -> str:
    """Chave de busca: sem espaços e em maiúsculo. 9036.0 (número no Excel) vira '9036'."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return re.sub(r"\s+", "", str(valor)).upper()


# --------------------------------------------------------------- leitura Excel
def _abas_xlsb(conteudo: bytes):
    from pyxlsb import open_workbook

    with open_workbook(io.BytesIO(conteudo)) as wb:
        for nome in wb.sheets:
            with wb.get_sheet(nome) as aba:
                yield nome, [[c.v for c in linha] for linha in aba.rows()]


def _abas_xlsx(conteudo: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            yield ws.title, [list(linha) for linha in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _abas_xls(conteudo: bytes):
    import xlrd

    wb = xlrd.open_workbook(file_contents=conteudo)
    for sh in wb.sheets():
        linhas = [[(v if v != "" else None) for v in sh.row_values(i)] for i in range(sh.nrows)]
        yield sh.name, linhas


def _extrair_itens(linhas: list) -> dict:
    """Acha o cabeçalho (CÓDIGO + REVENDA) nas primeiras linhas e lê os itens abaixo dele."""
    cab = col_cod = col_rev = col_desc = None
    for i, linha in enumerate(linhas[:LINHAS_PARA_ACHAR_CABECALHO]):
        nomes = [_texto_chave(c) if c is not None else "" for c in linha]
        if "CODIGO" in nomes and "REVENDA" in nomes:
            cab, col_cod, col_rev = i, nomes.index("CODIGO"), nomes.index("REVENDA")
            col_desc = next((j for j, n in enumerate(nomes) if n.startswith("DESCRICAO")), None)
            break
    if cab is None:
        return {}

    itens = {}
    for linha in linhas[cab + 1:]:
        if col_cod >= len(linha) or col_rev >= len(linha):
            continue
        codigo = normalizar_codigo(linha[col_cod])
        revenda = linha[col_rev]
        if not codigo or isinstance(revenda, bool) or not isinstance(revenda, (int, float)):
            continue  # linha em branco, assinatura do rodapé, cabeçalho repetido etc.
        desc = linha[col_desc] if col_desc is not None and col_desc < len(linha) else None
        original = linha[col_cod]
        if isinstance(original, float) and original.is_integer():
            original = int(original)
        itens[codigo] = {
            "codigo": str(original).strip(),
            "descricao": " ".join(str(desc).split()) if desc else "",
            "revenda": round(float(revenda), 2),
        }
    return itens


# ------------------------------------------------------------------ persistência
def _salvar_no_disco() -> None:
    ARQUIVO_JSON.parent.mkdir(parents=True, exist_ok=True)
    ARQUIVO_JSON.write_text(json.dumps(_tabela, ensure_ascii=False), encoding="utf-8")


def _carregar_do_disco() -> None:
    global _tabela
    if ARQUIVO_JSON.exists():
        try:
            dados = json.loads(ARQUIVO_JSON.read_text(encoding="utf-8"))
            if isinstance(dados.get("itens"), dict):
                _tabela = dados
        except (json.JSONDecodeError, OSError):
            pass  # arquivo corrompido: segue sem tabela até subir outra planilha


_carregar_do_disco()


# -------------------------------------------------------------------- API pública
def status() -> dict:
    return {
        "carregada": bool(_tabela["itens"]),
        "arquivo": _tabela["arquivo"],
        "carregada_em": _tabela["carregada_em"],
        "total_itens": len(_tabela["itens"]),
    }


def carregar_planilha(nome_arquivo: str, conteudo: bytes) -> dict:
    """Lê a planilha, substitui a tabela atual e devolve o status. ValueError = mensagem pro usuário."""
    ext = Path(nome_arquivo or "").suffix.lower()
    if ext not in EXTENSOES:
        raise ValueError("Formato não suportado. Envie um arquivo .xlsx, .xlsb, .xlsm ou .xls.")
    leitor = {".xlsb": _abas_xlsb, ".xls": _abas_xls}.get(ext, _abas_xlsx)

    itens: dict = {}
    try:
        for nome_aba, linhas in leitor(conteudo):
            for chave, item in _extrair_itens(linhas).items():
                item["aba"] = nome_aba
                itens.setdefault(chave, item)  # código repetido: vale a primeira ocorrência
    except ImportError as e:
        raise ValueError(f"Falta instalar uma biblioteca no servidor ({e.name}). Rode o calculadora.bat de novo.")
    except Exception:
        raise ValueError("Não consegui ler esta planilha. Confira se o arquivo não está corrompido ou protegido por senha.")

    if not itens:
        raise ValueError("Não encontrei as colunas CÓDIGO e REVENDA (com valores numéricos) em nenhuma aba da planilha.")

    global _tabela
    with _lock:
        _tabela = {
            "arquivo": Path(nome_arquivo).name,
            "carregada_em": datetime.now().strftime("%d/%m/%Y %H:%M"),
            "itens": itens,
        }
        _salvar_no_disco()
    return status()


def buscar(codigo: str) -> dict | None:
    return _tabela["itens"].get(normalizar_codigo(codigo))


def sugerir(codigo: str, limite: int = 5) -> list[str]:
    """Códigos parecidos (começam com / contêm o digitado), pra ajudar quando não acha."""
    chave = normalizar_codigo(codigo)
    if len(chave) < 3:
        return []
    comecam = [v["codigo"] for k, v in _tabela["itens"].items() if k.startswith(chave)]
    contem = [v["codigo"] for k, v in _tabela["itens"].items() if chave in k and not k.startswith(chave)]
    return (comecam + contem)[:limite]


def avaliar(codigo: str | None, preco_promocional: float) -> dict | None:
    """Compara o preço promocional calculado com a REVENDA do PMA. None = nenhum código informado."""
    codigo = (codigo or "").strip()
    if not codigo:
        return None
    item = buscar(codigo)
    if item is None:
        return {"codigo": codigo, "encontrado": False, "tabela_carregada": bool(_tabela["itens"])}
    revenda = item["revenda"]
    abaixo = round(preco_promocional, 2) < revenda
    return {
        "codigo": item["codigo"],
        "encontrado": True,
        "tabela_carregada": True,
        "descricao": item["descricao"],
        "aba": item.get("aba"),
        "marca": marca_da_aba(item.get("aba")),
        "revenda": revenda,
        "abaixo": abaixo,
        "diferenca": round(revenda - preco_promocional, 2) if abaixo else None,
    }
