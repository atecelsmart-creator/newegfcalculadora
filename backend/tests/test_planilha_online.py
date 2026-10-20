"""Modo online da planilha: enviar -> localizar -> gravar -> ver -> baixar (Mercado Livre, xlsx e csv)."""
import io
import os
import time
import uuid

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api.planilha_online import router
from app.planilha_ml import online

app = FastAPI()
app.include_router(router, prefix="/api")


@pytest.fixture()
def cli(tmp_path, monkeypatch):
    monkeypatch.setattr(online, "PASTA", tmp_path / "sessoes")
    online._cache.clear()
    return TestClient(app)


def xlsx_ml(n=17) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Anúncios"
    ws.append(["ITEM_ID", "VARIATION_ID", "SKU", "TITLE", "PRICE", "SALE_PRICE", "STATUS"])
    ws.append(["Código do anúncio", "Variação", "SKU", "Título", "Preço", "Preço em promoção", "Status"])   # linha de ajuda
    for i in range(1, n + 1):
        ws.append([f"MLB{3705773800 + i}", None, f"SKU{i:03d}", f"Furadeira modelo {i}", 100.0 + i, None, "Ativo"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def enviar(cli, sessao, conteudo=None, nome="Anuncios-2026_10_04-16_25.xlsx", mk="mercado_livre"):
    return cli.post("/api/planilha-online/enviar", data={"sessao": sessao, "marketplace": mk},
                    files={"planilha": (nome, xlsx_ml() if conteudo is None else conteudo, "application/octet-stream")})


def test_fluxo_completo_mercado_livre(cli):
    s = str(uuid.uuid4())
    assert cli.post("/api/planilha-online/resumo", data={"sessao": s}).json() == {"carregado": False}

    r = enviar(cli, s)
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["total_anuncios"] == 17 and d["nome_download"] == "Anuncios-2026_10_04-16_25_atualizado.xlsx"

    # localizar por SKU e por código do anúncio
    a = cli.post("/api/planilha-online/localizar", data={"sessao": s, "busca": "SKU005"}).json()
    assert a["preco_atual"] == 105.0 and a["alterado"] is False
    assert cli.post("/api/planilha-online/localizar", data={"sessao": s, "busca": "MLB3705773805"}).json()["sku"] == "SKU005"

    # gravar duas vezes na mesma linha: o "original" continua sendo o da planilha enviada
    g = cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "SKU005", "preco_cheio": "199.9", "preco_promocional": "179.9"})
    assert g.status_code == 200, g.text
    assert g.json()["conferencia_ok"] and g.json()["gravou_promocional"] and g.json()["alterados"] == 1
    cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "SKU005", "preco_cheio": "210", "preco_promocional": "190"})
    cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "SKU009", "preco_cheio": "150"})

    # a aba "Planilha online" vê as alterações
    v = cli.post("/api/planilha-online/visualizar", data={"sessao": s}).json()
    assert v["total"] == 17 and v["alterados"] == 2 and len(v["linhas"]) == 17
    l5 = next(l for l in v["linhas"] if l["sku"] == "SKU005")
    assert (l5["preco_atual"], l5["promocao_atual"], l5["preco_original"], l5["alterado"]) == (210.0, 190.0, 105.0, True)
    # a tela marca as linhas alteradas pelo campo "alterado" (o filtro so_alterados do servidor não existe mais)
    assert [l["sku"] for l in v["linhas"] if l["alterado"]] == ["SKU005", "SKU009"]
    filtro = cli.post("/api/planilha-online/visualizar", data={"sessao": s, "busca": "modelo 1"}).json()
    assert filtro["filtrados"] == 9   # modelo 1, 10..17
    pag = cli.post("/api/planilha-online/visualizar", data={"sessao": s, "por_pagina": "10", "pagina": "2"}).json()
    assert pag["paginas"] == 2 and len(pag["linhas"]) == 7

    # baixar: é um xlsx válido, com os preços gravados e o resto intacto
    b = cli.post("/api/planilha-online/baixar", data={"sessao": s})
    assert b.status_code == 200 and "_atualizado.xlsx" in b.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(b.content))["Anúncios"]
    assert ws["E7"].value == 210.0 and ws["F7"].value == 190.0 and ws["E11"].value == 150.0
    assert ws["E3"].value == 101.0 and ws["D3"].value == "Furadeira modelo 1"

    # o estado está em disco: depois de limpar a memória, o resumo reconecta
    online._cache.clear()
    assert cli.post("/api/planilha-online/resumo", data={"sessao": s}).json()["alterados"] == 2

    assert cli.post("/api/planilha-online/descartar", data={"sessao": s}).json() == {"carregado": False}
    assert cli.post("/api/planilha-online/baixar", data={"sessao": s}).status_code == 422


def test_cor_das_atualizadas_e_limpar_uma_linha(cli):
    s = str(uuid.uuid4())
    enviar(cli, s)
    post = lambda rota, **d: cli.post(f"/api/planilha-online/{rota}", data={"sessao": s, **d})
    cor_de = lambda sku: next(l for l in post("visualizar").json()["linhas"] if l["sku"] == sku)

    # a linha gravada recebe a cor escolhida; sem cor escolhida, fica só "alterada" (verde padrão na tela)
    assert post("gravar", busca="SKU005", preco_cheio="199", cor="#fef08a").status_code == 200
    assert post("gravar", busca="SKU009", preco_cheio="150").status_code == 200
    assert (cor_de("SKU005")["cor"], cor_de("SKU005")["alterado"]) == ("#fef08a", True)
    assert (cor_de("SKU009")["cor"], cor_de("SKU009")["alterado"]) == ("", True)

    # gravar de novo com outra cor troca a cor da linha
    post("gravar", busca="SKU005", preco_cheio="205", cor="#bfdbfe")
    assert cor_de("SKU005")["cor"] == "#bfdbfe"

    # limpar UMA linha: só ela perde cor e marca; a outra continua e o preço não muda
    linha5 = cor_de("SKU005")["linha"]
    assert post("limpar-marcas", linhas=str(linha5)).status_code == 200
    assert (cor_de("SKU005")["cor"], cor_de("SKU005")["alterado"], cor_de("SKU005")["preco_atual"]) == ("", False, 205.0)
    assert cor_de("SKU009")["alterado"] is True

    # depois de limpar, uma nova gravação marca a linha de novo com a cor escolhida
    post("gravar", busca="SKU005", preco_cheio="210", cor="#fbcfe8")
    assert (cor_de("SKU005")["cor"], cor_de("SKU005")["alterado"]) == ("#fbcfe8", True)

    # cor inválida é recusada antes de gravar qualquer coisa
    assert post("gravar", busca="SKU001", preco_cheio="500", cor="vermelho").status_code == 422
    assert cor_de("SKU001")["preco_atual"] == 101.0


def test_sessoes_sao_isoladas(cli):
    s1, s2 = str(uuid.uuid4()), str(uuid.uuid4())
    enviar(cli, s1)
    cli.post("/api/planilha-online/gravar", data={"sessao": s1, "busca": "SKU001", "preco_cheio": "500"})
    assert cli.post("/api/planilha-online/resumo", data={"sessao": s2}).json() == {"carregado": False}
    assert cli.post("/api/planilha-online/localizar", data={"sessao": s2, "busca": "SKU001"}).status_code == 422


def test_validacoes(cli):
    s = str(uuid.uuid4())
    assert enviar(cli, "../../etc").status_code == 422                       # sessão inválida (path traversal)
    assert enviar(cli, s, nome="x.xls").status_code == 422                   # formato só do modo local
    assert enviar(cli, s, conteudo=b"").status_code == 422
    assert enviar(cli, s, conteudo=b"nao e planilha").status_code == 422
    assert enviar(cli, s, mk="shopee").status_code == 422                    # planilha do ML no marketplace errado
    assert cli.post("/api/planilha-online/localizar", data={"sessao": s, "busca": "X"}).status_code == 422   # nada enviado
    enviar(cli, s)
    assert cli.post("/api/planilha-online/localizar", data={"sessao": s, "busca": "NAOEXISTE"}).status_code == 422
    assert cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "SKU001", "preco_cheio": "-5"}).status_code == 422


def test_promocional_invalido_grava_so_o_cheio(cli):
    s = str(uuid.uuid4())
    enviar(cli, s)
    g = cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "SKU002", "preco_cheio": "100", "preco_promocional": "120"}).json()
    assert g["gravou_promocional"] is False and g["avisos"]


def test_csv(cli):
    s = str(uuid.uuid4())
    csv = "ITEM_ID;VARIATION_ID;SKU;TITLE;PRICE;SALE_PRICE;STATUS\nMLB1234567890;;ABC1;Serra;10,50;;Ativo\n"
    assert enviar(cli, s, conteudo=csv.encode(), nome="ml.csv").status_code == 200
    g = cli.post("/api/planilha-online/gravar", data={"sessao": s, "busca": "ABC1", "preco_cheio": "12.5", "preco_promocional": "11"})
    assert g.status_code == 200 and g.json()["conferencia_ok"]
    b = cli.post("/api/planilha-online/baixar", data={"sessao": s})
    assert b.content.decode().splitlines()[1] == "MLB1234567890;;ABC1;Serra;12,50;11,00;Ativo"
    assert b.headers["content-type"].startswith("text/csv")


def test_limpeza_de_arquivos_antigos(cli):
    velha = str(uuid.uuid4())
    enviar(cli, velha)
    for f in (online.PASTA / velha).rglob("*"):
        os.utime(f, (time.time() - 13 * 3600,) * 2)
    enviar(cli, str(uuid.uuid4()))                                          # um novo envio limpa as sessões velhas
    assert not (online.PASTA / velha).exists()
